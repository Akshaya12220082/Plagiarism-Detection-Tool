# utils/text_extractor.py
from docx import Document
import fitz
from pptx import Presentation

def extract_text_from_path(path):
    ext = path.rsplit('.', 1)[-1].lower()
    try:
        if ext == 'txt':
            with open(path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        elif ext == 'docx':
            doc = Document(path)
            return '\n'.join(p.text for p in doc.paragraphs)
        elif ext == 'pdf':
            text_parts = []
            with fitz.open(path) as pdf:
                for page in pdf:
                    text_parts.append(page.get_text())
            return '\n'.join(text_parts)
        elif ext == 'pptx':
            prs = Presentation(path)
            texts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        texts.append(shape.text)
            return '\n'.join(texts)
        else:
            return ''
    except Exception as e:
        print("extract_text error:", e)
        return ''
